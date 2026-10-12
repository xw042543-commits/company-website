from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


INSTITUTIONS = (
    (("亚太", "APU"), "apu"),
    (("英迪", "INTI"), "inti"),
    (("双威", "Sunway", "MUFY", "CIMP"), "sunway"),
    (("蒙纳士", "莫纳什", "Monash"), "monash"),
    (("泰莱", "Taylor"), "taylor"),
    (("思特雅", "UCSI"), "ucsi"),
    (("精英", "HELP"), "help"),
    (("世纪", "SEGi"), "segi"),
    (("南安普顿", "Southampton"), "southampton"),
    (("诺丁汉", "Nottingham"), "nottingham"),
    (("马大", "马来亚", "University of Malaya"), "um"),
)


def guess_institution(name: str) -> str:
    lowered = name.lower()
    for aliases, slug in INSTITUTIONS:
        if any(alias.lower() in lowered for alias in aliases):
            return slug
    return "unknown"


def guess_category(name: str) -> str:
    if re.search(r"宿舍|住宿|accommod|residence", name, re.I):
        return "accommodation"
    if re.search(r"签证", name, re.I):
        return "visa"
    if re.search(r"学费|fee", name, re.I):
        return "tuition"
    if re.search(r"语言|雅思|IELTS", name, re.I):
        return "language_requirements"
    if re.search(r"入学|录取|高考|Gaokao", name, re.I):
        return "admission_requirements"
    if re.search(r"校历|时间节点|intake", name, re.I):
        return "intakes"
    if re.search(r"课程|专业|本科|硕士|博士|预科|招生|手册|宣传", name, re.I):
        return "programme_catalogue"
    return "reference"


def guess_year(name: str) -> str:
    years = re.findall(r"20(?:2[4-9]|3\d)", name)
    if years:
        return max(years)
    match = re.search(r"(?<!\d)(2[4-9])(?!\d)", name)
    return f"20{match.group(1)}" if match else ""


def extract_pdf(path: Path, output: Path) -> tuple[int, int, str]:
    reader = PdfReader(path)
    blocks: list[str] = []
    errors: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text(extraction_mode="layout") or ""
        except Exception as exc:  # keep the batch moving and record the bad page
            text = ""
            errors.append(f"page {index}: {exc}")
        blocks.append(f"=== PAGE {index} ===\n{text.strip()}\n")
    joined = "\n".join(blocks)
    output.write_text(joined, encoding="utf-8")
    return len(reader.pages), len(joined), "; ".join(errors)


def extract_docx(path: Path, output: Path) -> tuple[int, int, str]:
    doc = Document(path)
    blocks = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    for table_index, table in enumerate(doc.tables, start=1):
        blocks.append(f"=== TABLE {table_index} ===")
        for row in table.rows:
            blocks.append("\t".join(cell.text.strip().replace("\n", " | ") for cell in row.cells))
    joined = "\n".join(blocks)
    output.write_text(joined, encoding="utf-8")
    return len(doc.paragraphs), len(joined), ""


def extract_xlsx(path: Path, output: Path) -> tuple[int, int, str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks: list[str] = []
    for sheet in workbook.worksheets:
        blocks.append(f"=== SHEET {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value).strip().replace("\n", " | ") for value in row]
            while values and not values[-1]:
                values.pop()
            if any(values):
                blocks.append("\t".join(values))
    joined = "\n".join(blocks)
    output.write_text(joined, encoding="utf-8")
    return len(workbook.sheetnames), len(joined), ""


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: extract_sources.py SOURCE_DIR OUTPUT_DIR")
    source_dir = Path(sys.argv[1])
    output_dir = Path(sys.argv[2])
    raw_dir = output_dir / "raw-text"
    raw_dir.mkdir(parents=True, exist_ok=True)

    rows: list[dict[str, object]] = []
    for path in sorted((item for item in source_dir.iterdir() if item.is_file()), key=lambda p: p.name):
        suffix = path.suffix.lower()
        unit_count = 0
        char_count = 0
        error = ""
        if suffix in {".pdf", ".docx", ".xlsx"}:
            output = raw_dir / f"{path.stem}.txt"
            try:
                if suffix == ".pdf":
                    unit_count, char_count, error = extract_pdf(path, output)
                elif suffix == ".docx":
                    unit_count, char_count, error = extract_docx(path, output)
                else:
                    unit_count, char_count, error = extract_xlsx(path, output)
            except Exception as exc:
                error = str(exc)
        rows.append(
            {
                "filename": path.name,
                "extension": suffix,
                "size_bytes": path.stat().st_size,
                "institution_slug": guess_institution(path.name),
                "category": guess_category(path.name),
                "source_year": guess_year(path.name),
                "units": unit_count,
                "text_chars": char_count,
                "needs_ocr": suffix in {".jpg", ".jpeg", ".png"} or (suffix == ".pdf" and char_count < 100 * max(unit_count, 1)),
                "error": error,
            }
        )

    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "source-manifest.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    with (output_dir / "source-manifest.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    print(json.dumps({"files": len(rows), "errors": sum(bool(row["error"]) for row in rows), "needs_ocr": sum(bool(row["needs_ocr"]) for row in rows)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
